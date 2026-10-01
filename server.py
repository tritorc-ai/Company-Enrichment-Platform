import os
import re
import io
import csv
import json
import time
import asyncio
from urllib.parse import urljoin, urlparse, quote_plus

import requests
import urllib3
import openpyxl
from bs4 import BeautifulSoup
from dotenv import load_dotenv
from fastapi import FastAPI, Request, UploadFile, File
from fastapi.responses import StreamingResponse, FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from groq import Groq

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

APP_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(APP_DIR, ".env"))
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) CompanyEnrichmentBot/1.0"
GROQ_MODEL = os.environ.get("GROQ_MODEL", "openai/gpt-oss-120b")

URL_LIKE_RE = re.compile(
    r"^(https?://)?([a-z0-9-]+\.)+[a-z]{2,}(/.*)?$", re.IGNORECASE
)
EMAIL_RE = re.compile(r"^[^\s@]+@([a-z0-9-]+(?:\.[a-z0-9-]+)+)$", re.IGNORECASE)

# consumer/webmail domains: an email at one of these tells us nothing about
# the sender's employer, so it must not be treated as a company website.
FREE_EMAIL_DOMAINS = {
    "gmail.com", "yahoo.com", "outlook.com", "hotmail.com", "live.com",
    "icloud.com", "aol.com", "gmx.com", "protonmail.com", "mail.com",
    "yandex.com", "qq.com", "163.com", "126.com",
}

with open(os.path.join(APP_DIR, "tritorc_offerings.json"), encoding="utf-8") as f:
    TRITORC_OFFERINGS = json.load(f)

# category-level summary (not per-product) to keep the grounding payload small
# enough to stay well under Groq's per-minute token limits across back-to-back calls.
TRITORC_OFFERINGS_COMPACT = {
    "product_categories": TRITORC_OFFERINGS["product_categories"],
    "services": TRITORC_OFFERINGS["services"],
}
TRITORC_BLOCK = json.dumps(TRITORC_OFFERINGS_COMPACT, ensure_ascii=False)

app = FastAPI(title="Company Enrichment Tool")


def get_groq_client():
    api_key = os.environ.get("GROQ_API_KEY")
    if not api_key:
        raise RuntimeError("GROQ_API_KEY environment variable is not set.")
    return Groq(api_key=api_key)


def _probe_url(url: str, method: str = "head"):
    """Reachability check only (no data exchanged). Some real corporate sites
    ship an incomplete/misconfigured TLS certificate chain (verified case:
    bilfinger.com fails standard cert validation while resolving and serving
    fine over plain TLS) — retry once without strict verification rather than
    silently treating the whole company as having 'no website'."""
    fn = requests.head if method == "head" else requests.get
    try:
        return fn(url, headers={"User-Agent": UA}, timeout=10, allow_redirects=True)
    except requests.exceptions.SSLError:
        try:
            return fn(url, headers={"User-Agent": UA}, timeout=10, allow_redirects=True, verify=False)
        except Exception:
            return None
    except Exception:
        return None


SOCIAL_AND_AGGREGATOR_HOSTS = {
    "linkedin.com", "www.linkedin.com", "en.wikipedia.org", "wikipedia.org",
    "x.com", "twitter.com", "facebook.com", "www.facebook.com",
    "instagram.com", "youtube.com", "www.youtube.com",
    "finance.yahoo.com", "www.crunchbase.com", "crunchbase.com",
    "www.bloomberg.com", "bloomberg.com", "www.glassdoor.com", "glassdoor.com",
    "indeed.com", "www.indeed.com",
}


def search_serper(company_name: str):
    """Live Google search via serper.dev. Returns the first organic result
    whose domain isn't a social/aggregator site, or None if unavailable/no match."""
    api_key = os.environ.get("SERPER_API_KEY")
    if not api_key:
        return None
    try:
        r = requests.post(
            "https://google.serper.dev/search",
            headers={"X-API-KEY": api_key, "Content-Type": "application/json"},
            json={"q": f"{company_name} official website", "num": 8},
            timeout=15,
        )
        if r.status_code != 200:
            return None
        data = r.json()
        for item in data.get("organic", []):
            link = item.get("link")
            if not link:
                continue
            host = urlparse(link).netloc.lower()
            if host in SOCIAL_AND_AGGREGATOR_HOSTS or "jobs." in host:
                continue
            parsed = urlparse(link)
            return f"{parsed.scheme}://{parsed.netloc}"
        return None
    except Exception:
        return None


def find_official_site(client, company_name: str):
    """Resolve a likely official website. Primary path is a live Google search
    via serper.dev (when SERPER_API_KEY is set), taking the first non-social,
    non-aggregator organic result. Falls back to asking the LLM for its
    best-known domain and verifying that guess actually loads before trusting
    it, if no search key is configured or the search finds nothing usable.
    (Live scraping of DuckDuckGo/Bing HTML was tried first and found
    unreliable — both returned bot-detection junk/decoy results inconsistently
    for automated queries.)
    """
    searched = search_serper(company_name)
    if searched:
        r = _probe_url(searched, method="get")
        if r is not None and r.status_code < 400:
            return str(r.url).rstrip("/")

    try:
        completion = client.chat.completions.create(
            model=GROQ_MODEL,
            messages=[
                {"role": "system", "content": "You return only valid JSON."},
                {"role": "user", "content": (
                    f'What is the primary official corporate website domain for the company "{company_name}"? '
                    'Respond with a JSON object: {"domain": "example.com"} using only the bare domain '
                    '(no scheme, no path). If you are not confident you know the real domain, '
                    'respond with {"domain": null}.'
                )},
            ],
            temperature=0,
            max_tokens=300,
            response_format={"type": "json_object"},
            reasoning_effort="low",
        )
        raw = completion.choices[0].message.content
        data = json.loads(raw)
        domain = data.get("domain")
        if not domain:
            return None
        domain = domain.strip().lower()
        domain = re.sub(r"^https?://", "", domain).split("/")[0]
        if not URL_LIKE_RE.match(domain):
            return None
        # try the bare domain, then www.-prefixed (some sites 404 on bare apex
        # and only serve from www., e.g. bilfinger.com vs www.bilfinger.com)
        candidates = [f"https://{domain}"]
        if not domain.startswith("www."):
            candidates.append(f"https://www.{domain}")

        for candidate in candidates:
            r = _probe_url(candidate)
            if r is not None and r.status_code >= 400:
                r = _probe_url(candidate, method="get")
            if r is not None and r.status_code < 400:
                return str(r.url).rstrip("/")
        return None
    except Exception:
        return None


CANDIDATE_PATH_KEYWORDS = [
    "about", "company", "services", "products", "solutions", "projects",
    "portfolio", "news", "media", "press", "industries", "capabilities",
]


def crawl_company_site(base_url: str, max_pages: int = 4, timeout: int = 15):
    """Light crawl: homepage + a handful of relevant internal links."""
    session = requests.Session()
    session.headers.update({"User-Agent": UA})
    visited = set()
    pages_text = []

    def clean(soup):
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        return re.sub(r"\n{2,}", "\n", soup.get_text("\n", strip=True))

    def session_get(url):
        # some real sites ship an incomplete TLS cert chain; retry once
        # without strict verification rather than dropping the whole crawl.
        try:
            return session.get(url, timeout=timeout)
        except requests.exceptions.SSLError:
            return session.get(url, timeout=timeout, verify=False)
        except requests.exceptions.RequestException:
            # verified case: some hosts time out on a direct HTTPS connection
            # but work fine when reached via their own http->https redirect
            # (e.g. sceptre.com.my). Retry over plain http as a last resort.
            if url.startswith("https://"):
                return session.get("http://" + url[len("https://"):], timeout=timeout)
            raise

    try:
        r = session_get(base_url)
    except Exception as e:
        return [], f"error fetching homepage: {e}"

    if r.status_code != 200:
        return [], f"homepage returned status {r.status_code}"

    soup = BeautifulSoup(r.text, "lxml")
    home_text = clean(soup)
    pages_text.append({"url": base_url, "text": home_text[:3000]})
    visited.add(base_url.rstrip("/"))

    candidates = []
    for a in soup.find_all("a", href=True):
        href = a["href"]
        full = urljoin(base_url, href)
        parsed = urlparse(full)
        if parsed.netloc != urlparse(base_url).netloc:
            continue
        path_lower = parsed.path.lower()
        if any(kw in path_lower for kw in CANDIDATE_PATH_KEYWORDS):
            norm = full.split("#")[0].rstrip("/")
            if norm not in visited:
                candidates.append(norm)

    seen_candidates = []
    for c in candidates:
        if c not in seen_candidates:
            seen_candidates.append(c)

    for url in seen_candidates[: max_pages - 1]:
        try:
            rr = session_get(url)
            time.sleep(0.5)
            if rr.status_code == 200:
                s2 = BeautifulSoup(rr.text, "lxml")
                pages_text.append({"url": url, "text": clean(s2)[:2000]})
                visited.add(url)
        except Exception:
            continue

    return pages_text, None


ENRICHMENT_SCHEMA_HINT = """{
  "company_name": "",
  "website": "",
  "country": "",
  "industry": "",
  "company_category": "",
  "business_description": "",
  "key_operations": [],
  "projects_or_recent_activity": [],
  "tritorc_relevance": []
}"""


def build_prompt(company_name, website, crawled_pages, no_site_found, extra_note=None):
    if no_site_found:
        source_block = extra_note or (
            "NO WEBSITE COULD BE FOUND OR CRAWLED. You may use general knowledge if you are confident about this company, "
            "but you MUST set every field you are not confident about to null or [], and set \"business_description\" "
            "to include a note that no live website source was available."
        )
    else:
        joined = "\n\n---PAGE BREAK---\n\n".join(
            f"URL: {p['url']}\n{p['text']}" for p in crawled_pages
        )
        source_block = f"WEBSITE CONTENT (scraped just now):\n{joined[:6000]}"

    return f"""You are a B2B research analyst. Research the company "{company_name}" using ONLY the source material given below (website text if provided). Do not invent facts not supported by the source; if something is unknown, use null or an empty list.

{source_block}

TASK: Return a single JSON object with EXACTLY this shape (no extra keys, no markdown fences, no commentary):
{ENRICHMENT_SCHEMA_HINT}

Field rules:
- "company_name": the company's proper name.
- "website": the root URL used as source, or null if none.
- "country": the company's primary country of operation/headquarters, or null if not stated.
- "industry": the company's primary industry (e.g. "Oil & Gas", "Power Generation", "Steel Manufacturing", "Wind Energy", "Mining", "Construction", "Petrochemical", etc.) based on the source text.
- "company_category": MUST be exactly one of "distributor", "ECP", "end_user". Use "distributor" if the company resells/distributes industrial tools or equipment. Use "ECP" if it is an engineering/construction/procurement contractor delivering projects for others. Use "end_user" for every other private company or government/public-sector organization that would use industrial tools/services in its own operations (this covers both private firms and government bodies).
- "business_description": 2-4 factual sentences describing what the company does, grounded in the source text.
- "key_operations": list of concrete operational activities/business lines mentioned in the source (e.g. "pipeline construction", "refinery operations", "turbine maintenance").
- "projects_or_recent_activity": list of specific named projects, plants, contracts, expansions, or news items mentioned in the source text. Empty list if none mentioned.
- "tritorc_relevance": list of short strings, each naming a SPECIFIC Tritorc product category, example product, or service from the catalog below AND why it's relevant to this company's operations (e.g. "Hydraulic Torque Wrenches (e.g. TSL Series) — relevant for flange bolting during the refinery turnarounds mentioned on their site"). Only reference items that actually appear in the catalog below. If nothing in the source material suggests a real need, return an empty list rather than forcing a match.

TRITORC PRODUCT & SERVICE CATALOG (only reference items from this list in tritorc_relevance):
{TRITORC_BLOCK}

Return ONLY the JSON object."""


def normalize_if_url(entry: str):
    """If the input line is already a URL/domain, return a normalized https:// URL. Else None."""
    candidate = entry.strip()
    if " " in candidate:
        return None
    if not URL_LIKE_RE.match(candidate):
        return None
    if not candidate.startswith("http://") and not candidate.startswith("https://"):
        candidate = "https://" + candidate
    return candidate


def call_groq_with_retry(client, prompt: str, max_retries: int = 4):
    from groq import APIStatusError

    last_err = None
    for attempt in range(max_retries):
        try:
            return client.chat.completions.create(
                model=GROQ_MODEL,
                messages=[
                    {"role": "system", "content": "You return only valid JSON. Never wrap it in markdown code fences."},
                    {"role": "user", "content": prompt},
                ],
                temperature=0.2,
                max_tokens=1500,
                response_format={"type": "json_object"},
                reasoning_effort="low",
            )
        except APIStatusError as e:
            last_err = e
            if e.status_code == 429:
                wait_s = 5.0
                try:
                    body = e.response.json()
                    msg = body.get("error", {}).get("message", "")
                    m = re.search(r"try again in ([\d.]+)s", msg)
                    if m:
                        wait_s = float(m.group(1)) + 0.5
                except Exception:
                    pass
                time.sleep(min(wait_s, 30))
                continue
            if e.status_code == 400 and "json_validate_failed" in str(e):
                # transient generation glitch (malformed/truncated JSON); retry immediately
                continue
            raise
    raise last_err


def resolve_from_email(email: str):
    """If the input is an email address, derive the company's website from its
    domain. Returns (website, display_name, domain, is_free_email_domain)."""
    m = EMAIL_RE.match(email.strip())
    if not m:
        return None
    domain = m.group(1).lower()
    if domain in FREE_EMAIL_DOMAINS:
        return (None, email.strip(), domain, True)
    return (f"https://{domain}", domain, domain, False)


def enrich_company(client, company_input: str):
    email_match = resolve_from_email(company_input)
    direct_url = normalize_if_url(company_input)
    is_free_email = False
    email_domain = None

    if email_match:
        website, display_name, email_domain, is_free_email = email_match
        if is_free_email:
            # a personal/webmail address carries no employer signal; don't
            # guess a company from it — the free-form enrichment path below
            # will honestly report "no website source available".
            website = None
    elif direct_url:
        website = direct_url
        display_name = urlparse(direct_url).netloc
    else:
        website = find_official_site(client, company_input)
        display_name = company_input

    pages, err = ([], "no website found") if not website else crawl_company_site(website)
    if not pages and website:
        # direct URL might have a trailing path or a different scheme; retry with bare origin
        parsed = urlparse(website)
        fallback = f"{parsed.scheme}://{parsed.netloc}"
        if fallback != website:
            pages, err = crawl_company_site(fallback)
            if pages:
                website = fallback
    no_site_found = not pages
    extra_note = (
        f'The input "{company_input}" is a personal/webmail email address (domain "{email_domain}" is a consumer '
        "email provider, not a company domain). It carries NO reliable signal about the sender's employer. "
        "Set every field to null or [] except business_description, which must state that the employer could not be "
        "determined from a personal email domain."
        if is_free_email else None
    )

    prompt = build_prompt(display_name, website, pages, no_site_found, extra_note=extra_note)
    completion = call_groq_with_retry(client, prompt)
    raw = completion.choices[0].message.content
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        data = json.loads(match.group(0)) if match else {
            "company_name": display_name, "website": website, "country": None,
            "industry": None, "company_category": None, "business_description": None,
            "key_operations": [], "projects_or_recent_activity": [], "tritorc_relevance": [],
        }

    data.setdefault("company_name", display_name)
    if not data.get("website"):
        data["website"] = website
    return data


@app.get("/")
async def index():
    return FileResponse(os.path.join(APP_DIR, "static", "index.html"))


app.mount("/static", StaticFiles(directory=os.path.join(APP_DIR, "static")), name="static")


NAME_KEY_CANDIDATES = ["company_name", "company", "name", "companies", "email", "email_address"]
HEADER_NAME_CANDIDATES = ("company_name", "company name", "company", "name", "email", "email address")


def _extract_names_from_json(data):
    if isinstance(data, list):
        names = []
        for item in data:
            if isinstance(item, str):
                names.append(item)
            elif isinstance(item, dict):
                for k in NAME_KEY_CANDIDATES:
                    if item.get(k):
                        names.append(str(item[k]))
                        break
        return names
    if isinstance(data, dict):
        for k in NAME_KEY_CANDIDATES:
            if isinstance(data.get(k), list):
                return _extract_names_from_json(data[k])
        return []
    return []


def _pick_name_column(header_row):
    """Given a header row, find the column index most likely to hold company
    names. Falls back to the first column if nothing matches."""
    if not header_row:
        return 0
    lowered = [str(h or "").strip().lower() for h in header_row]
    for candidate in ["company_name", "company name", "company", "name", "organization", "organisation", "email", "email address"]:
        if candidate in lowered:
            return lowered.index(candidate)
    return 0


def parse_companies_file(filename: str, content: bytes):
    """Extract a flat list of company names from an uploaded .txt/.csv/.xlsx/.json file."""
    ext = os.path.splitext(filename.lower())[1]

    if ext == ".json":
        data = json.loads(content.decode("utf-8"))
        return [n.strip() for n in _extract_names_from_json(data) if str(n).strip()]

    if ext == ".csv":
        text = content.decode("utf-8-sig", errors="ignore")
        rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            return []
        col = _pick_name_column(rows[0])
        looks_like_header = _pick_name_column(rows[0]) != 0 or any(
            str(h or "").strip().lower() in HEADER_NAME_CANDIDATES
            for h in rows[0]
        )
        data_rows = rows[1:] if looks_like_header else rows
        return [r[col].strip() for r in data_rows if len(r) > col and r[col].strip()]

    if ext in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            return []
        col = _pick_name_column(rows[0])
        looks_like_header = any(
            str(h or "").strip().lower() in HEADER_NAME_CANDIDATES
            for h in rows[0]
        )
        data_rows = rows[1:] if looks_like_header else rows
        return [str(r[col]).strip() for r in data_rows if len(r) > col and r[col] and str(r[col]).strip()]

    # .txt and anything else: treat as plain text, one name per line (or comma-separated)
    text = content.decode("utf-8", errors="ignore")
    parts = re.split(r"[\r\n,]+", text)
    return [p.strip() for p in parts if p.strip()]


@app.post("/api/parse-file")
async def api_parse_file(file: UploadFile = File(...)):
    content = await file.read()
    try:
        names = parse_companies_file(file.filename, content)
    except Exception as e:
        return JSONResponse({"error": f"Could not parse {file.filename}: {e}"}, status_code=400)
    return {"companies": names}


EXPORT_COLUMNS = [
    "company_name", "website", "country", "industry", "company_category",
    "business_description", "key_operations", "projects_or_recent_activity",
    "tritorc_relevance",
]
LIST_COLUMNS = {"key_operations", "projects_or_recent_activity", "tritorc_relevance"}


def _row_for(result: dict):
    row = []
    for col in EXPORT_COLUMNS:
        val = result.get(col)
        if col in LIST_COLUMNS:
            row.append("; ".join(val) if isinstance(val, list) else (val or ""))
        else:
            row.append(val if val is not None else "")
    return row


@app.post("/api/export-xlsx")
async def api_export_xlsx(request: Request):
    body = await request.json()
    results = body.get("results", [])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Company Enrichment"
    header = [c.replace("_", " ").title() for c in EXPORT_COLUMNS]
    ws.append(header)
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    for result in results:
        ws.append(_row_for(result))

    widths = [22, 26, 14, 20, 16, 50, 40, 40, 50]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=company_enrichment.xlsx"},
    )


@app.post("/api/enrich")
async def api_enrich(request: Request):
    body = await request.json()
    names = body.get("companies", [])
    names = [n.strip() for n in names if n.strip()]

    async def event_stream():
        try:
            client = get_groq_client()
        except RuntimeError as e:
            yield f"data: {json.dumps({'type': 'error', 'message': str(e)})}\n\n"
            return

        results = []
        total = len(names)
        yield f"data: {json.dumps({'type': 'start', 'total': total})}\n\n"

        loop = asyncio.get_event_loop()
        for idx, name in enumerate(names, start=1):
            yield f"data: {json.dumps({'type': 'progress', 'index': idx, 'total': total, 'company': name, 'status': 'processing'})}\n\n"
            await asyncio.sleep(0)
            try:
                data = await loop.run_in_executor(None, enrich_company, client, name)
            except Exception as e:
                data = {
                    "company_name": name, "website": None, "country": None, "industry": None,
                    "company_category": None, "business_description": f"Enrichment failed: {e}",
                    "key_operations": [], "projects_or_recent_activity": [], "tritorc_relevance": [],
                }
            results.append(data)
            yield f"data: {json.dumps({'type': 'progress', 'index': idx, 'total': total, 'company': name, 'status': 'done'})}\n\n"

        yield f"data: {json.dumps({'type': 'complete', 'results': results})}\n\n"

    return StreamingResponse(event_stream(), media_type="text/event-stream")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
